import json
import os
import shutil
import threading
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal, Optional

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

load_dotenv()

# ---------------------------------------------------------------------------
# Trava de segurança: no Render, DATA_DIR é obrigatório (mesmo padrão dos
# outros projetos Piaseg) — sem isso, um deploy sem a env var configurada
# gravaria dados no container efêmero e perderia tudo silenciosamente.
# ---------------------------------------------------------------------------
if os.getenv("RENDER") and not os.getenv("DATA_DIR"):
    raise RuntimeError(
        "DATA_DIR não está configurado no Render. Isso faria o app gravar dados "
        "fora do disco persistente, perdendo tudo no próximo deploy. Configure a "
        "env var DATA_DIR=/data no serviço antes de tentar de novo."
    )

app = FastAPI(title="Piaseg Desconto Comercial - Conta Corrente")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

APP_DIR = Path(__file__).parent
DATA_DIR = Path(os.getenv("DATA_DIR", str(APP_DIR / "data_local")))
DATA_DIR.mkdir(parents=True, exist_ok=True)

SEGURADORAS_FILE = DATA_DIR / "seguradoras.json"
LANCAMENTOS_FILE = DATA_DIR / "lancamentos.json"
ADMIN_FILE = DATA_DIR / "admin.json"
BACKUPS_DIR = DATA_DIR / "backups"
MAX_BACKUPS = 60
# Além do snapshot a cada deploy, tira um snapshot automático antes de uma
# gravação se o último tiver mais de 1h — assim um erro de operação no meio
# do dia também tem ponto de volta, não só os deploys.
AUTO_SNAPSHOT_INTERVAL_S = 3600
GESTORES_FILE = DATA_DIR / "gestores.json"
SOLICITACOES_FILE = DATA_DIR / "solicitacoes.json"
_DATA_FILES = ["seguradoras.json", "lancamentos.json", "gestores.json", "solicitacoes.json"]
_BACKUP_FILES = _DATA_FILES + ["admin.json"]
LOGOS_DIR = DATA_DIR / "logos"
LOGO_MAX_BYTES = 2 * 1024 * 1024
LOGO_TIPOS = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp", "image/svg+xml": "svg"}

# Lista de franqueados para o formulário de solicitação vem do Piaseg
# Franqueados (só os nomes dos ativos), com cache para não depender dele a
# cada acesso; se estiver fora do ar o formulário aceita digitação livre.
FRANQUEADOS_API_URL = os.getenv("FRANQUEADOS_API_URL", "https://piaseg-franqueados-backend.onrender.com/franqueados")
FRANQUEADOS_CACHE_S = 600

ADMIN_PASSWORD_SEED = os.getenv("ADMIN_PASSWORD", "")

# Tipos de lançamento:
#  credito    — entra saldo de desconto na conta da seguradora
#  utilizacao — desconto usado direto (sai do saldo)
#  reserva    — valor separado para um uso previsto; enquanto "aberta" reduz
#               o "saldo a utilizar"; ao ser baixada vira utilização
#               (status "utilizada"); se cancelada, libera o valor.
TIPOS = {"credito", "utilizacao", "reserva"}

_lock = threading.Lock()


def _load(path: Path) -> list:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return []


def _save(path: Path, data) -> None:
    _maybe_auto_snapshot()
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)  # gravação atômica: nunca deixa o arquivo pela metade


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _snapshot(prefix: str = "") -> str:
    name = prefix + datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%S")
    snapshot_dir = BACKUPS_DIR / name
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    for filename in _BACKUP_FILES:
        source = DATA_DIR / filename
        if source.exists():
            shutil.copy2(source, snapshot_dir / filename)
    if LOGOS_DIR.exists():
        shutil.copytree(LOGOS_DIR, snapshot_dir / "logos", dirs_exist_ok=True)
    existing = sorted((p for p in BACKUPS_DIR.iterdir() if p.is_dir()), key=lambda p: p.stat().st_mtime)
    for old in existing[:-MAX_BACKUPS]:
        shutil.rmtree(old, ignore_errors=True)
    return name


def _maybe_auto_snapshot() -> None:
    if not os.getenv("DATA_DIR") or not BACKUPS_DIR.exists():
        return
    dirs = [p for p in BACKUPS_DIR.iterdir() if p.is_dir()]
    latest = max((p.stat().st_mtime for p in dirs), default=0)
    if datetime.now(timezone.utc).timestamp() - latest > AUTO_SNAPSHOT_INTERVAL_S:
        _snapshot("auto-")


def _run_startup_backup() -> None:
    """Snapshot dos dados a cada subida do app (= cada deploy). Só roda com
    DATA_DIR explicitamente setado (produção), nunca em dev local."""
    if not os.getenv("DATA_DIR"):
        return
    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    _snapshot("deploy-")


_run_startup_backup()


# ---------------------------------------------------------------------------
# Auth admin (senha única compartilhada; cada operador informa o nome, que
# fica gravado em todo lançamento que ele fizer)
# ---------------------------------------------------------------------------

bearer = HTTPBearer()


def _load_admin_password() -> str:
    if not ADMIN_FILE.exists():
        ADMIN_FILE.write_text(json.dumps({"password": ADMIN_PASSWORD_SEED}), encoding="utf-8")
        return ADMIN_PASSWORD_SEED
    return json.loads(ADMIN_FILE.read_text(encoding="utf-8")).get("password", "")


def require_admin(credentials: HTTPAuthorizationCredentials = Depends(bearer)) -> None:
    senha_atual = _load_admin_password()
    if not senha_atual or credentials.credentials != senha_atual:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Senha inválida")


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------


class AdminLogin(BaseModel):
    password: str


class SenhaUpdate(BaseModel):
    senha_nova: str


class SeguradoraIn(BaseModel):
    nome: str
    ativa: bool = True


class LancamentoIn(BaseModel):
    tipo: Literal["credito", "utilizacao", "reserva"]
    valor: float = Field(gt=0)
    descricao: str = ""
    franqueado: str = ""
    data: Optional[str] = None  # AAAA-MM-DD; padrão = hoje
    operador: str


class AcaoOperador(BaseModel):
    operador: str
    motivo: str = ""


class GestorIn(BaseModel):
    nome: str
    email: str = ""
    ativo: bool = True


class SolicitacaoIn(BaseModel):
    seguradora_id: str
    valor: float = Field(gt=0)
    franqueado: str
    gestor_id: str
    observacao: str = ""


class RespostaSolicitacao(BaseModel):
    decisao: Literal["aprovar", "reprovar"]
    resposta: str = ""
    operador: str


# ---------------------------------------------------------------------------
# Cálculo de saldos
# ---------------------------------------------------------------------------


def _cents(valor: float) -> int:
    return int(round(valor * 100))


def _logo_url(seguradora: dict) -> Optional[str]:
    # ?v= muda a cada upload, para o navegador não mostrar a logo antiga do cache
    logo = seguradora.get("logo")
    return f'/logos/{logo["arquivo"]}?v={logo["versao"]}' if logo else None


def _resumo(seguradora: dict, lancamentos: list[dict]) -> dict:
    creditos = utilizado = reservado = 0
    ultimo = None
    for l in lancamentos:
        if l["seguradora_id"] != seguradora["id"]:
            continue
        ultimo = max(ultimo or "", l.get("updated_at") or l["created_at"])
        if l.get("estornado"):
            continue
        if l["tipo"] == "credito":
            creditos += l["valor_centavos"]
        elif l["tipo"] == "utilizacao":
            utilizado += l["valor_centavos"]
        elif l["tipo"] == "reserva":
            if l.get("status") == "aberta":
                reservado += l["valor_centavos"]
            elif l.get("status") == "utilizada":
                utilizado += l["valor_centavos"]
    saldo = creditos - utilizado
    return {
        "id": seguradora["id"],
        "nome": seguradora["nome"],
        "logo_url": _logo_url(seguradora),
        "ativa": seguradora.get("ativa", True),
        "creditos": creditos / 100,
        "utilizado": utilizado / 100,
        "saldo": saldo / 100,
        "reserva": reservado / 100,
        "saldo_a_utilizar": (saldo - reservado) / 100,
        "atualizado_em": ultimo or seguradora.get("created_at"),
    }


def _seguradora_or_404(seguradoras: list[dict], seg_id: str) -> dict:
    for s in seguradoras:
        if s["id"] == seg_id:
            return s
    raise HTTPException(404, "Seguradora não encontrada")


def _lancamento_or_404(lancamentos: list[dict], lanc_id: str) -> dict:
    for l in lancamentos:
        if l["id"] == lanc_id:
            return l
    raise HTTPException(404, "Lançamento não encontrado")


def _public_lancamento(l: dict) -> dict:
    out = dict(l)
    out["valor"] = l["valor_centavos"] / 100
    return out


def _operador(nome: str) -> str:
    nome = (nome or "").strip()
    if not nome:
        raise HTTPException(400, "Informe o nome do operador")
    return nome


# ---------------------------------------------------------------------------
# Rotas públicas
# ---------------------------------------------------------------------------


@app.get("/health")
def health():
    return {"ok": True}


@app.get("/painel")
def painel():
    """Painel aberto aos franqueados: só os números por seguradora ativa."""
    seguradoras = _load(SEGURADORAS_FILE)
    lancamentos = _load(LANCAMENTOS_FILE)
    linhas = [_resumo(s, lancamentos) for s in seguradoras if s.get("ativa", True)]
    linhas.sort(key=lambda r: r["nome"].lower())
    publicas = [
        {k: r[k] for k in ("id", "nome", "logo_url", "saldo", "reserva", "saldo_a_utilizar", "atualizado_em")}
        for r in linhas
    ]
    return {
        "seguradoras": publicas,
        "totais": {
            "saldo": round(sum(r["saldo"] for r in linhas), 2),
            "reserva": round(sum(r["reserva"] for r in linhas), 2),
            "saldo_a_utilizar": round(sum(r["saldo_a_utilizar"] for r in linhas), 2),
        },
        "atualizado_em": max((r["atualizado_em"] or "" for r in linhas), default=None) or None,
    }


# ---------------------------------------------------------------------------
# Admin
# ---------------------------------------------------------------------------


@app.post("/admin/login")
def admin_login(body: AdminLogin):
    senha = _load_admin_password()
    if not senha or body.password != senha:
        raise HTTPException(401, "Senha inválida")
    return {"ok": True}


@app.put("/admin/senha")
def trocar_senha(body: SenhaUpdate, _: None = Depends(require_admin)):
    if len(body.senha_nova) < 6:
        raise HTTPException(400, "A senha precisa ter pelo menos 6 caracteres")
    with _lock:
        _save(ADMIN_FILE, {"password": body.senha_nova})
    return {"ok": True}


@app.get("/admin/seguradoras")
def list_seguradoras(_: None = Depends(require_admin)):
    seguradoras = _load(SEGURADORAS_FILE)
    lancamentos = _load(LANCAMENTOS_FILE)
    linhas = [_resumo(s, lancamentos) for s in seguradoras]
    linhas.sort(key=lambda r: (not r["ativa"], r["nome"].lower()))
    return linhas


@app.post("/admin/seguradoras", status_code=201)
def create_seguradora(body: SeguradoraIn, _: None = Depends(require_admin)):
    nome = body.nome.strip()
    if not nome:
        raise HTTPException(400, "Informe o nome da seguradora")
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        if any(s["nome"].strip().lower() == nome.lower() for s in seguradoras):
            raise HTTPException(409, "Já existe uma seguradora com esse nome")
        seg = {"id": uuid.uuid4().hex[:10], "nome": nome, "ativa": body.ativa, "created_at": now_iso()}
        seguradoras.append(seg)
        _save(SEGURADORAS_FILE, seguradoras)
    return seg


@app.put("/admin/seguradoras/{seg_id}")
def update_seguradora(seg_id: str, body: SeguradoraIn, _: None = Depends(require_admin)):
    nome = body.nome.strip()
    if not nome:
        raise HTTPException(400, "Informe o nome da seguradora")
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        seg = _seguradora_or_404(seguradoras, seg_id)
        if any(s["id"] != seg_id and s["nome"].strip().lower() == nome.lower() for s in seguradoras):
            raise HTTPException(409, "Já existe uma seguradora com esse nome")
        seg["nome"] = nome
        seg["ativa"] = body.ativa
        _save(SEGURADORAS_FILE, seguradoras)
    return seg


@app.delete("/admin/seguradoras/{seg_id}")
def delete_seguradora(seg_id: str, _: None = Depends(require_admin)):
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        _seguradora_or_404(seguradoras, seg_id)
        if any(l["seguradora_id"] == seg_id for l in _load(LANCAMENTOS_FILE)):
            raise HTTPException(
                409, "Essa seguradora já tem lançamentos. Para tirar do painel, desative em vez de excluir."
            )
        _save(SEGURADORAS_FILE, [s for s in seguradoras if s["id"] != seg_id])
    return {"ok": True}


@app.post("/admin/seguradoras/{seg_id}/logo")
async def upload_logo(seg_id: str, arquivo: UploadFile = File(...), _: None = Depends(require_admin)):
    ext = LOGO_TIPOS.get(arquivo.content_type or "")
    if not ext:
        raise HTTPException(400, "Formato não aceito. Use PNG, JPG, WEBP ou SVG.")
    conteudo = await arquivo.read()
    if len(conteudo) > LOGO_MAX_BYTES:
        raise HTTPException(400, "A imagem pode ter no máximo 2 MB.")
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        seg = _seguradora_or_404(seguradoras, seg_id)
        LOGOS_DIR.mkdir(parents=True, exist_ok=True)
        nome_arquivo = f"{seg_id}.{ext}"
        for antigo in LOGOS_DIR.glob(f"{seg_id}.*"):
            if antigo.name != nome_arquivo:
                antigo.unlink()
        (LOGOS_DIR / nome_arquivo).write_bytes(conteudo)
        seg["logo"] = {"arquivo": nome_arquivo, "versao": uuid.uuid4().hex[:8]}
        _save(SEGURADORAS_FILE, seguradoras)
    return {"logo_url": _logo_url(seg)}


@app.delete("/admin/seguradoras/{seg_id}/logo")
def remove_logo(seg_id: str, _: None = Depends(require_admin)):
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        seg = _seguradora_or_404(seguradoras, seg_id)
        seg.pop("logo", None)  # arquivo fica no disco (backups); só deixa de aparecer
        _save(SEGURADORAS_FILE, seguradoras)
    return {"ok": True}


@app.get("/logos/{arquivo}")
def get_logo(arquivo: str):
    caminho = LOGOS_DIR / arquivo
    if "/" in arquivo or ".." in arquivo or not caminho.is_file():
        raise HTTPException(404, "Logo não encontrada")
    ext = caminho.suffix.lstrip(".")
    media = {v: k for k, v in LOGO_TIPOS.items()}.get(ext, "application/octet-stream")
    return FileResponse(
        caminho,
        media_type=media,
        headers={
            "Cache-Control": "public, max-age=31536000",
            # SVG pode carregar script: bloqueia qualquer execução se abrirem a URL direto
            "Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.get("/admin/seguradoras/{seg_id}")
def get_seguradora(seg_id: str, _: None = Depends(require_admin)):
    seguradoras = _load(SEGURADORAS_FILE)
    seg = _seguradora_or_404(seguradoras, seg_id)
    lancamentos = _load(LANCAMENTOS_FILE)
    extrato = [_public_lancamento(l) for l in lancamentos if l["seguradora_id"] == seg_id]
    extrato.sort(key=lambda l: (l["data"], l["created_at"]), reverse=True)
    return {"resumo": _resumo(seg, lancamentos), "lancamentos": extrato}


@app.post("/admin/seguradoras/{seg_id}/lancamentos", status_code=201)
def create_lancamento(seg_id: str, body: LancamentoIn, _: None = Depends(require_admin)):
    operador = _operador(body.operador)
    data = body.data or datetime.now(timezone.utc).date().isoformat()
    try:
        datetime.strptime(data, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(400, "Data inválida")
    with _lock:
        _seguradora_or_404(_load(SEGURADORAS_FILE), seg_id)
        lancamentos = _load(LANCAMENTOS_FILE)
        lanc = {
            "id": uuid.uuid4().hex[:12],
            "seguradora_id": seg_id,
            "tipo": body.tipo,
            "valor_centavos": _cents(body.valor),
            "descricao": body.descricao.strip(),
            "franqueado": body.franqueado.strip(),
            "data": data,
            "operador": operador,
            "created_at": now_iso(),
            "estornado": False,
        }
        if body.tipo == "reserva":
            lanc["status"] = "aberta"
        lancamentos.append(lanc)
        _save(LANCAMENTOS_FILE, lancamentos)
    return _public_lancamento(lanc)


def _acao_reserva(lanc_id: str, body: AcaoOperador, novo_status: str) -> dict:
    operador = _operador(body.operador)
    with _lock:
        lancamentos = _load(LANCAMENTOS_FILE)
        lanc = _lancamento_or_404(lancamentos, lanc_id)
        if lanc["tipo"] != "reserva" or lanc.get("status") != "aberta" or lanc.get("estornado"):
            raise HTTPException(400, "Só é possível baixar ou cancelar uma reserva em aberto")
        lanc["status"] = novo_status
        lanc["status_por"] = operador
        lanc["status_em"] = now_iso()
        lanc["status_motivo"] = body.motivo.strip()
        lanc["updated_at"] = lanc["status_em"]
        _save(LANCAMENTOS_FILE, lancamentos)
    return _public_lancamento(lanc)


@app.post("/admin/lancamentos/{lanc_id}/utilizar")
def utilizar_reserva(lanc_id: str, body: AcaoOperador, _: None = Depends(require_admin)):
    return _acao_reserva(lanc_id, body, "utilizada")


@app.post("/admin/lancamentos/{lanc_id}/cancelar")
def cancelar_reserva(lanc_id: str, body: AcaoOperador, _: None = Depends(require_admin)):
    return _acao_reserva(lanc_id, body, "cancelada")


@app.post("/admin/lancamentos/{lanc_id}/estornar")
def estornar_lancamento(lanc_id: str, body: AcaoOperador, _: None = Depends(require_admin)):
    """Anula um lançamento feito por engano. Nada é apagado: o lançamento
    continua no extrato, marcado como estornado, com quem fez e o motivo."""
    operador = _operador(body.operador)
    if not body.motivo.strip():
        raise HTTPException(400, "Informe o motivo do estorno")
    with _lock:
        lancamentos = _load(LANCAMENTOS_FILE)
        lanc = _lancamento_or_404(lancamentos, lanc_id)
        if lanc.get("estornado"):
            raise HTTPException(400, "Esse lançamento já foi estornado")
        lanc["estornado"] = True
        lanc["estornado_por"] = operador
        lanc["estornado_em"] = now_iso()
        lanc["estorno_motivo"] = body.motivo.strip()
        lanc["updated_at"] = lanc["estornado_em"]
        _save(LANCAMENTOS_FILE, lancamentos)
    return _public_lancamento(lanc)


# ---------------------------------------------------------------------------
# Gestores (assessores comerciais que recebem as solicitações)
# ---------------------------------------------------------------------------


@app.get("/gestores")
def list_gestores_publico():
    gestores = [g for g in _load(GESTORES_FILE) if g.get("ativo", True)]
    gestores.sort(key=lambda g: g["nome"].lower())
    return [{"id": g["id"], "nome": g["nome"]} for g in gestores]


@app.get("/admin/gestores")
def list_gestores(_: None = Depends(require_admin)):
    gestores = _load(GESTORES_FILE)
    gestores.sort(key=lambda g: (not g.get("ativo", True), g["nome"].lower()))
    return gestores


@app.post("/admin/gestores", status_code=201)
def create_gestor(body: GestorIn, _: None = Depends(require_admin)):
    nome = body.nome.strip()
    if not nome:
        raise HTTPException(400, "Informe o nome do gestor")
    with _lock:
        gestores = _load(GESTORES_FILE)
        if any(g["nome"].strip().lower() == nome.lower() for g in gestores):
            raise HTTPException(409, "Já existe um gestor com esse nome")
        g = {"id": uuid.uuid4().hex[:10], "nome": nome, "email": body.email.strip(), "ativo": body.ativo, "created_at": now_iso()}
        gestores.append(g)
        _save(GESTORES_FILE, gestores)
    return g


@app.put("/admin/gestores/{gestor_id}")
def update_gestor(gestor_id: str, body: GestorIn, _: None = Depends(require_admin)):
    nome = body.nome.strip()
    if not nome:
        raise HTTPException(400, "Informe o nome do gestor")
    with _lock:
        gestores = _load(GESTORES_FILE)
        g = next((x for x in gestores if x["id"] == gestor_id), None)
        if not g:
            raise HTTPException(404, "Gestor não encontrado")
        g.update({"nome": nome, "email": body.email.strip(), "ativo": body.ativo})
        _save(GESTORES_FILE, gestores)
    return g


@app.delete("/admin/gestores/{gestor_id}")
def delete_gestor(gestor_id: str, _: None = Depends(require_admin)):
    with _lock:
        if any(s["gestor_id"] == gestor_id for s in _load(SOLICITACOES_FILE)):
            raise HTTPException(409, "Esse gestor já recebeu solicitações. Desative em vez de excluir.")
        gestores = _load(GESTORES_FILE)
        if not any(g["id"] == gestor_id for g in gestores):
            raise HTTPException(404, "Gestor não encontrado")
        _save(GESTORES_FILE, [g for g in gestores if g["id"] != gestor_id])
    return {"ok": True}


_franq_cache: dict = {"em": 0.0, "nomes": []}


@app.get("/franqueados-lista")
def franqueados_lista():
    import urllib.request

    agora = datetime.now(timezone.utc).timestamp()
    if agora - _franq_cache["em"] > FRANQUEADOS_CACHE_S or not _franq_cache["nomes"]:
        try:
            with urllib.request.urlopen(FRANQUEADOS_API_URL, timeout=8) as r:
                dados = json.loads(r.read().decode("utf-8"))
            nomes = sorted(
                {(f.get("nome_fantasia") or "").strip() for f in dados if f.get("status", "ativo") == "ativo"} - {""},
                key=str.lower,
            )
            _franq_cache.update({"em": agora, "nomes": nomes})
        except Exception as e:  # fora do ar: devolve o último cache (ou vazio)
            print(f"[franqueados] falha ao buscar lista: {e}")
    return _franq_cache["nomes"]


# ---------------------------------------------------------------------------
# Solicitações de reserva (franqueado pede → gestor aprova ou não)
# ---------------------------------------------------------------------------

SOLIC_STATUS_LABEL = {"pendente": "Aguardando análise", "aprovada": "Aprovada", "reprovada": "Não aprovada"}


def _solicitacao_view(s: dict, seguradoras: list[dict], gestores: list[dict], lancamentos: list[dict]) -> dict:
    seg = next((x for x in seguradoras if x["id"] == s["seguradora_id"]), None)
    gest = next((x for x in gestores if x["id"] == s["gestor_id"]), None)
    out = {k: v for k, v in s.items() if k != "valor_centavos"}
    out["valor"] = s["valor_centavos"] / 100
    out["seguradora_nome"] = seg["nome"] if seg else "?"
    out["gestor_nome"] = gest["nome"] if gest else "?"
    if s.get("lancamento_id"):
        lanc = next((l for l in lancamentos if l["id"] == s["lancamento_id"]), None)
        if lanc:
            out["reserva_status"] = "estornada" if lanc.get("estornado") else lanc.get("status")
    return out


@app.post("/solicitacoes", status_code=201)
def create_solicitacao(body: SolicitacaoIn):
    franqueado = body.franqueado.strip()
    if not franqueado:
        raise HTTPException(400, "Informe o franqueado")
    with _lock:
        seguradoras = _load(SEGURADORAS_FILE)
        seg = next((s for s in seguradoras if s["id"] == body.seguradora_id and s.get("ativa", True)), None)
        if not seg:
            raise HTTPException(400, "Seguradora inválida")
        gestores = _load(GESTORES_FILE)
        if not any(g["id"] == body.gestor_id and g.get("ativo", True) for g in gestores):
            raise HTTPException(400, "Selecione o gestor que vai analisar a solicitação")
        solicitacoes = _load(SOLICITACOES_FILE)
        s = {
            # id longo e aleatório: é ele que dá acesso à página de acompanhamento
            "id": uuid.uuid4().hex,
            "numero": max((x["numero"] for x in solicitacoes), default=0) + 1,
            "seguradora_id": seg["id"],
            "valor_centavos": _cents(body.valor),
            "franqueado": franqueado,
            "gestor_id": body.gestor_id,
            "observacao": body.observacao.strip(),
            "status": "pendente",
            "created_at": now_iso(),
        }
        solicitacoes.append(s)
        _save(SOLICITACOES_FILE, solicitacoes)
    return _solicitacao_view(s, seguradoras, gestores, [])


@app.get("/solicitacoes/{sol_id}")
def get_solicitacao(sol_id: str):
    s = next((x for x in _load(SOLICITACOES_FILE) if x["id"] == sol_id), None)
    if not s:
        raise HTTPException(404, "Solicitação não encontrada")
    return _solicitacao_view(s, _load(SEGURADORAS_FILE), _load(GESTORES_FILE), _load(LANCAMENTOS_FILE))


@app.get("/admin/solicitacoes")
def list_solicitacoes(_: None = Depends(require_admin)):
    seguradoras, gestores, lancamentos = _load(SEGURADORAS_FILE), _load(GESTORES_FILE), _load(LANCAMENTOS_FILE)
    saldos = {s["id"]: _resumo(s, lancamentos)["saldo_a_utilizar"] for s in seguradoras}
    out = []
    for s in sorted(_load(SOLICITACOES_FILE), key=lambda x: x["numero"], reverse=True):
        v = _solicitacao_view(s, seguradoras, gestores, lancamentos)
        v["saldo_a_utilizar_atual"] = saldos.get(s["seguradora_id"], 0)
        out.append(v)
    return out


@app.post("/admin/solicitacoes/{sol_id}/responder")
def responder_solicitacao(sol_id: str, body: RespostaSolicitacao, _: None = Depends(require_admin)):
    operador = _operador(body.operador)
    resposta = body.resposta.strip()
    if body.decisao == "reprovar" and not resposta:
        raise HTTPException(400, "Informe a justificativa para o franqueado")
    with _lock:
        solicitacoes = _load(SOLICITACOES_FILE)
        s = next((x for x in solicitacoes if x["id"] == sol_id), None)
        if not s:
            raise HTTPException(404, "Solicitação não encontrada")
        if s["status"] != "pendente":
            raise HTTPException(400, "Essa solicitação já foi respondida")
        agora = now_iso()
        if body.decisao == "aprovar":
            # aprovação já cria a reserva na conta corrente da seguradora
            lancamentos = _load(LANCAMENTOS_FILE)
            descricao = f'Solicitação #{s["numero"]}'
            if s.get("observacao"):
                descricao += f' — {s["observacao"]}'
            lanc = {
                "id": uuid.uuid4().hex[:12],
                "seguradora_id": s["seguradora_id"],
                "tipo": "reserva",
                "status": "aberta",
                "valor_centavos": s["valor_centavos"],
                "descricao": descricao,
                "franqueado": s["franqueado"],
                "data": datetime.now(timezone.utc).date().isoformat(),
                "operador": operador,
                "created_at": agora,
                "estornado": False,
                "solicitacao_id": s["id"],
            }
            lancamentos.append(lanc)
            _save(LANCAMENTOS_FILE, lancamentos)
            s["lancamento_id"] = lanc["id"]
        s["status"] = "aprovada" if body.decisao == "aprovar" else "reprovada"
        s["resposta"] = resposta
        s["respondido_por"] = operador
        s["respondido_em"] = agora
        _save(SOLICITACOES_FILE, solicitacoes)
    return _solicitacao_view(s, _load(SEGURADORAS_FILE), _load(GESTORES_FILE), _load(LANCAMENTOS_FILE))


TIPO_LABEL = {"credito": "Crédito", "utilizacao": "Utilização", "reserva": "Reserva"}
STATUS_LABEL = {"aberta": "Em aberto", "utilizada": "Utilizada", "cancelada": "Cancelada"}


def _br_date(iso: str) -> str:
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return iso


@app.get("/admin/export")
def export_excel(_: None = Depends(require_admin)):
    seguradoras = _load(SEGURADORAS_FILE)
    lancamentos = _load(LANCAMENTOS_FILE)
    nomes = {s["id"]: s["nome"] for s in seguradoras}

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="072A3C")
    brl = '"R$" #,##0.00;[Red]-"R$" #,##0.00'

    ws = wb.active
    ws.title = "Saldos"
    ws.append(["Seguradora", "Situação", "Saldo", "Reserva", "Saldo à utilizar"])
    for r in sorted((_resumo(s, lancamentos) for s in seguradoras), key=lambda r: r["nome"].lower()):
        ws.append([r["nome"], "Ativa" if r["ativa"] else "Inativa", r["saldo"], r["reserva"], r["saldo_a_utilizar"]])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            c.number_format = brl

    ws2 = wb.create_sheet("Extrato")
    ws2.append(["Data", "Seguradora", "Tipo", "Valor", "Situação", "Franqueado", "Descrição", "Lançado por", "Observações"])
    for l in sorted(lancamentos, key=lambda l: (l["data"], l["created_at"])):
        if l.get("estornado"):
            situacao = "Estornado"
            obs = f'Estornado por {l.get("estornado_por", "")}: {l.get("estorno_motivo", "")}'
        elif l["tipo"] == "reserva":
            situacao = STATUS_LABEL.get(l.get("status", ""), "")
            obs = f'{situacao} por {l["status_por"]}' if l.get("status_por") else ""
            if l.get("status_motivo"):
                obs += f': {l["status_motivo"]}'
        else:
            situacao, obs = "", ""
        ws2.append([
            _br_date(l["data"]), nomes.get(l["seguradora_id"], "?"), TIPO_LABEL[l["tipo"]],
            l["valor_centavos"] / 100, situacao, l.get("franqueado", ""), l.get("descricao", ""),
            l.get("operador", ""), obs,
        ])
    for row in ws2.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = brl

    for sheet in (ws, ws2):
        for c in sheet[1]:
            c.font = header_font
            c.fill = header_fill
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 50)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"desconto-comercial-{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Backups
# ---------------------------------------------------------------------------


@app.get("/admin/backups")
def list_backups(_: None = Depends(require_admin)):
    if not BACKUPS_DIR.exists():
        return []
    dirs = sorted((p for p in BACKUPS_DIR.iterdir() if p.is_dir()), key=lambda p: p.stat().st_mtime, reverse=True)
    return [
        {"name": p.name, "criado_em": datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat()}
        for p in dirs
    ]


@app.post("/admin/backups/{name}/restore")
def restore_backup(name: str, _: None = Depends(require_admin)):
    snapshot_dir = BACKUPS_DIR / name
    if "/" in name or ".." in name or not snapshot_dir.is_dir():
        raise HTTPException(404, "Backup não encontrado")
    with _lock:
        seguranca = _snapshot("antes-restore-")  # restore errado também tem volta
        for filename in _DATA_FILES:
            src = snapshot_dir / filename
            if src.exists():
                shutil.copy2(src, DATA_DIR / filename)
        if (snapshot_dir / "logos").is_dir():
            shutil.copytree(snapshot_dir / "logos", LOGOS_DIR, dirs_exist_ok=True)
    return {"ok": True, "backup_seguranca": seguranca}
